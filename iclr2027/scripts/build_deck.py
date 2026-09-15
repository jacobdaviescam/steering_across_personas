import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
S = os.path.dirname(os.path.abspath(__file__)); F = os.path.join(S, "figs")
src = Presentation(os.path.join(S, "iclr2027_results_2026-09-15.pptx"))
W, H = src.slide_width, src.slide_height
# slide index (1-based) -> figure file; tables on these slides move to an appendix
FIG = {6: "fig1_gemma_contexts.png", 7: "fig1b_llama_contexts.png", 8: "fig1c_qwen_contexts.png", 9: "fig3_axis.png", 10: "fig4_olmo_stages.png", 11: "fig5_nonsense_variants.png"}
moved = []  # (title, table element)
slides = list(src.slides)
for idx, slide in enumerate(slides, 1):
    if idx not in FIG: continue
    title = slide.shapes.title.text if slide.shapes.title else next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    tbl = [sh for sh in slide.shapes if sh.shape_type == 19]
    for t in tbl:
        moved.append((title, copy.deepcopy(t._element))); t._element.getparent().remove(t._element)
    # shrink remaining text boxes into a right-hand column? Simpler: keep title at top, push notes text to bottom strip
    texts = [sh for sh in slide.shapes if sh.has_text_frame and sh != slide.shapes.title]
    # first text box is the title (TextBox 1); others are notes/footer
    boxes = [sh for sh in slide.shapes if sh.has_text_frame]
    title_box, others = boxes[0], boxes[1:]
    top = title_box.top + title_box.height + Emu(60000)
    foot_h = Inches(1.15) if others else 0
    pic_h = H - top - foot_h - Emu(120000)
    pic = slide.shapes.add_picture(os.path.join(F, FIG[idx]), Inches(0.4), top, height=pic_h)
    if pic.width > W - Inches(0.8):
        ratio = (W - Inches(0.8)) / pic.width; pic.width = W - Inches(0.8); pic.height = int(pic.height * ratio)
    pic.left = int((W - pic.width) / 2)
    y = pic.top + pic.height + Emu(60000)
    for sh in others:
        sh.top = y; sh.left = Inches(0.4); sh.width = W - Inches(0.8); sh.height = Inches(0.5)
        for para in sh.text_frame.paragraphs:
            for r in para.runs: r.font.size = Pt(10)
        y += Inches(0.5)
# insert cross-model forest slide after slide 8
blank = src.slide_layouts[6]
def add_fig_slide(title, fig, footer):
    s = src.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.3), W - Inches(0.8), Inches(0.8)); tf = tb.text_frame; tf.word_wrap = True
    tf.text = title; tf.paragraphs[0].runs[0].font.size = Pt(22); tf.paragraphs[0].runs[0].font.bold = True
    pic = s.shapes.add_picture(os.path.join(F, fig), Inches(0.4), Inches(1.2), width=W - Inches(0.8))
    if pic.top + pic.height > H - Inches(0.9):
        r = (H - Inches(0.9) - pic.top) / pic.height; pic.height = int(pic.height * r); pic.width = int(pic.width * r); pic.left = int((W - pic.width) / 2)
    fb = s.shapes.add_textbox(Inches(0.4), H - Inches(0.8), W - Inches(0.8), Inches(0.5)); fb.text_frame.word_wrap = True
    fb.text_frame.text = footer; fb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    return s
cross = add_fig_slide("Cross-model: the automated grader splits Llama, the coding harness splits Gemma, nothing resolves on Qwen", "fig2_paired_deltas.png",
                      "Context minus nonsense, cosine to null, 95% paired question-bootstrap intervals (200 redraws). Label = traits whose interval lies entirely below zero.")
# move the new slide to position 9 (after Qwen slide)
sldIdLst = src.slides._sldIdLst; ids = list(sldIdLst); new = ids[-1]; sldIdLst.remove(new); sldIdLst.insert(8, new)
# appendix D: the tables that were on result slides
for title, el in moved:
    s = src.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.3), W - Inches(0.8), Inches(0.8)); tb.text_frame.word_wrap = True
    tb.text_frame.text = "Appendix D — table for: " + title; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(16); tb.text_frame.paragraphs[0].runs[0].font.bold = True
    s.shapes._spTree.append(el)
    # place the table below the heading
    gf = s.shapes[-1]; gf.top = Inches(1.3); gf.left = Inches(0.4)
out = os.path.join(S, "iclr2027_results_2026-09-15_v2.pptx"); src.save(out)
print("saved", out, "slides:", len(src.slides))
