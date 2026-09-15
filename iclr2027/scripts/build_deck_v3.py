"""Deck v3: start from v2, swap figures, rewrite the Qwen slide, add the OLMo-by-stage slide, refresh summary/next-steps."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
S = os.path.dirname(os.path.abspath(__file__)); F = os.path.join(S, "figs")
p = Presentation(os.path.join(S, "iclr2027_results_2026-09-15_v2.pptx")); W, H = p.slide_width, p.slide_height
slides = list(p.slides)

def boxes(s): return [sh for sh in s.shapes if sh.has_text_frame]
def set_text(sh, text, size=None, bold=None):
    tf = sh.text_frame; tf.word_wrap = True
    first = tf.paragraphs[0]; font = first.runs[0].font if first.runs else None
    sz = size or (font.size if font else None); bd = bold if bold is not None else (font.bold if font else None)
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    lines = text.split("\n"); first.text = lines[0]
    for r in first.runs: r.font.size = sz; r.font.bold = bd
    for ln in lines[1:]:
        para = tf.add_paragraph(); para.text = ln
        for r in para.runs: r.font.size = sz; r.font.bold = False
def swap_pic(s, fig):
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    if not pics: return
    old = pics[0]; left, top, width, height = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    pic = s.shapes.add_picture(os.path.join(F, fig), left, top, height=height)
    if pic.width > W - Inches(0.8):
        r = (W - Inches(0.8)) / pic.width; pic.width = W - Inches(0.8); pic.height = int(pic.height * r)
    pic.left = int((W - pic.width) / 2)

# 6 Gemma, 9 forest, 10 axis, 11 OLMo shared variance, 12 band
swap_pic(slides[5], "fig1_gemma_contexts.png"); swap_pic(slides[9], "fig3_axis.png"); swap_pic(slides[10], "fig4_olmo_stages.png")
# 8 Qwen slide -> regimes
s = slides[7]; swap_pic(s, "fig1c_qwen_regimes.png"); b = boxes(s)
set_text(b[0], "Qwen3: the gibberish floor breaks only through the chat template; under a fixed format base and instruct behave alike")
set_text(b[-1], "Qwen3-8B, layer 18 of 36, same seven contexts. Through the chat template the five-prompt band drops to 0.76 to 0.94 and no context clears it; under the raw Context/Question/Answer format the band sits at 0.94 to 0.95, floors at 0.99, and the grader and human-evaluator contexts fall below it on both base and instruct. Qwen3-32B (L31/L42) and Qwen3-14B reproduce this; Qwen2.5-32B under its template is intermediate.", size=Pt(10))
# 9 forest
s = slides[8]; swap_pic(s, "fig2_paired_deltas.png"); b = boxes(s)
set_text(b[0], "Cross-model: under a fixed prompt format every family shows the split; the chat template is what breaks Qwen3's floor")
set_text(b[-1], "Context minus the single gibberish prompt, cosine to null, 95% paired question-bootstrap intervals (200 redraws). Label = traits whose interval lies entirely below zero. Panels labelled by prompt regime.", size=Pt(10))
# 12 band
s = slides[11]; swap_pic(s, "fig5_nonsense_band.png"); b = boxes(s)
set_text(b[0], "The gibberish floor is a band, and the chat template widens it: measure the floor in the same regime as the contexts")
# new slide: OLMo contexts by stage, inserted after slide 11
blank = p.slide_layouts[6]
def add_fig_slide(title, fig, footer, pos):
    s = p.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.3), W - Inches(0.8), Inches(0.9)); tf = tb.text_frame; tf.word_wrap = True
    tf.text = title; tf.paragraphs[0].runs[0].font.size = Pt(20); tf.paragraphs[0].runs[0].font.bold = True
    pic = s.shapes.add_picture(os.path.join(F, fig), Inches(0.4), Inches(1.3), width=W - Inches(0.8))
    if pic.top + pic.height > H - Inches(1.0):
        r = (H - Inches(1.0) - pic.top) / pic.height; pic.height = int(pic.height * r); pic.width = int(pic.width * r); pic.left = int((W - pic.width) / 2)
    fb = s.shapes.add_textbox(Inches(0.4), H - Inches(0.9), W - Inches(0.8), Inches(0.6)); fb.text_frame.word_wrap = True
    fb.text_frame.text = footer; fb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    lst = p.slides._sldIdLst; ids = list(lst); new = ids[-1]; lst.remove(new); lst.insert(pos, new)
    return s
add_fig_slide("OLMo-2 stages: post-training broadens which contexts reshape trait directions; DPO and RLVR add breadth, not depth", "fig6_olmo_context_stages.png",
              "Raw format on every stage, layer 15. Grader and human-evaluator contexts already rotate the base model; SFT sharpens them and adds deployment; DPO and RLVR bring simulation, harness and observation contexts past the control (contexts clearing it: 2 at base, 7 at Instruct). The control itself drifts toward null, 0.950 to 0.961.", pos=11)
# summary slide (2) and next steps (13 -> now 14)
s = slides[1]; b = boxes(s)
set_text(b[0], "Summary: the operational situation reshapes trait directions; the chat template is itself a context; post-training broadens conditionality")
body = [
 "• Gemma-2-27B: coding harness rotates confidence to 0.70 against null, persona scale; automated grader below the gibberish control in 7/8 traits; who-is-watching contexts an order of magnitude smaller; unobserved = gibberish.",
 "• Trait rotation is orthogonal to the assistant axis: projecting the axis out changes cosine by at most 0.006 in 144 cells; evaluation contexts leave the model on-axis.",
 "• OLMo-2: three quarters of the reported SFT drop was the prompt-format switch; the real decline is 0.968 to 0.929, monotone through RLVR.",
 "• Llama-3.1-8B: the automated grader splits from the human evaluator in 7/8 traits (risk-taking -0.13 to -0.18).",
 "• Qwen3 (8B, 14B, 32B): the low gibberish floor appears only through the chat template; under a fixed raw format base and instruct match and show the grader/human split. Qwen2.5-32B is intermediate.",
 "• OLMo-2 by stage: grader/human rotation exists at base, SFT sharpens it, DPO and RLVR widen it to simulation, harness and observation contexts (2 to 7 contexts clearing the control).",
 "• Honesty is the most context-invariant direction on every model; risk-taking and confidence the most sensitive.",
]
set_text(b[1], "\n".join(body), size=Pt(12))
s = slides[12]; b = boxes(s)
set_text(b[0], "Done, deferred, dropped")
set_text(b[1], "\n".join([
 "• Done: Result C on Gemma 2, Llama, Qwen3-8B/14B/32B (template and raw), Qwen2.5-32B; assistant axis; OLMo raw-format control; OLMo evaluation contexts by stage; base-vs-instruct pairs at 8B and 14B.",
 "• Deferred: gpt-oss-120b reward-hacking RL trajectory (Tinker LoRA cannot be applied to fused MXFP4 experts outside Tinker; needs merged bf16 checkpoints or Tinker-side runs).",
 "• Dropped for this submission: Gemma-3-27B-IT evaluation contexts; persona-adapter zoo subset.",
 "• Next: abstract 18 Sep; paper 25 Sep; results on branch iclr2027/results and the private HF dataset jacobdavies/iclr2027-conditionality-results.",
]), size=Pt(12))
# appendix table titles that referenced the old Qwen claim
for s in slides[16:]:
    b = boxes(s)
    if b and "Qwen3-32B: no evaluation-context effect" in b[0].text_frame.text:
        set_text(b[0], "Appendix D — table for: Qwen3-32B under its chat template (see the regime slide for the raw-format result)")
out = os.path.join(S, "iclr2027_results_2026-09-15_v3.pptx"); p.save(out); print("saved", out, "slides:", len(p.slides))
